import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

output_docx_path = r"d:\SIP_N_BITE\Documentation\Sip_N_Bite_Complete_Project_Documentation.docx"
root_copy_path = r"d:\SIP_N_BITE\Sip_N_Bite_Complete_Project_Documentation.docx"

img_base_ppt = r"d:\SIP_N_BITE\Documentation\PPT\extracted_screens_full"
img_base_diva = r"d:\SIP_N_BITE\Documentation\PPT\extracted_screens_divA"
img_base_gen = r"d:\SIP_N_BITE\Documentation\final_doc\generated_diagrams"
img_base_media = r"d:\SIP_N_BITE\Documentation\final_doc\extracted_final_media"

cache_png_dir = r"d:\SIP_N_BITE\Documentation\temp_data\cache_png"
os.makedirs(cache_png_dir, exist_ok=True)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_table(table, header_bg="1A365D", alt_bg="F7FAFC", col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CBD5E0", sz="4")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for cell in hdr_cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for row_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for cell in row.cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(45, 55, 72)
                    
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)

def add_callout_box(doc, text, title="KEY ARCHITECTURAL HIGHLIGHT:"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EBF8FF")
    set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2B6CB0"/>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_title = p.add_run(f"{title} ")
    r_title.font.name = 'Calibri'
    r_title.font.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(43, 108, 176)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(9)
    r_text.font.italic = True
    r_text.font.color.rgb = RGBColor(45, 55, 72)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def ensure_compatible_png(img_path):
    """Converts any image to clean, standardized RGBA/RGB PNG for 100% python-docx compatibility."""
    if not os.path.exists(img_path):
        return None
    try:
        base_name = os.path.splitext(os.path.basename(img_path))[0] + ".png"
        target_path = os.path.join(cache_png_dir, base_name)
        with Image.open(img_path) as im:
            if im.mode in ('RGBA', 'LA') or (im.mode == 'P' and 'transparency' in im.info):
                im_rgba = im.convert('RGBA')
                im_rgba.save(target_path, 'PNG')
            else:
                im_rgb = im.convert('RGB')
                im_rgb.save(target_path, 'PNG')
        return target_path
    except Exception as e:
        print(f"Error converting image {img_path}: {e}")
        return img_path

def add_fig_image(doc, img_path, caption, width_in=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Image not found at {img_path}")
        return
    
    safe_path = ensure_compatible_png(img_path)
    if not safe_path or not os.path.exists(safe_path):
        safe_path = img_path
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(safe_path, width=Inches(width_in))
    
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_before = Pt(2)
    cap_p.paragraph_format.space_after = Pt(12)
    r_cap = cap_p.add_run(caption)
    r_cap.font.name = 'Calibri'
    r_cap.font.size = Pt(8.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(74, 85, 104)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(26, 54, 93)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(43, 108, 176)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_p(doc, text, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_bold = p.add_run(bold_prefix)
    r_bold.font.name = 'Calibri'
    r_bold.font.bold = True
    r_bold.font.size = Pt(10)
    r_bold.font.color.rgb = RGBColor(26, 54, 93)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = RGBColor(45, 55, 72)
    return p

print("builder_foundation.py updated with robust image handling.")
